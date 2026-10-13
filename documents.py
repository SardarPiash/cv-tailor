"""One layout/content stream feeds both DOCX and PDF renderers."""
from io import BytesIO
from datetime import date, datetime
from html import escape
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader


def date_label(value):
    if not value:
        return ''
    if value.casefold() in ('present', 'current', 'now'):
        return 'Present'
    for fmt in ('%Y-%m-%d', '%Y-%m', '%b %Y', '%B %Y', '%Y'):
        try:
            parsed = datetime.strptime(value, fmt)
            return parsed.strftime('%Y' if fmt == '%Y' else '%b %Y')
        except ValueError:
            pass
    return value


def date_range(start, end):
    return ' - '.join(filter(None, (date_label(start), date_label(end))))


def contact(profile):
    return ' | '.join(str(profile[k]) for k in ('email', 'phone', 'location', 'linkedin', 'portfolio') if profile.get(k))


def cv_blocks(profile, cv):
    blocks = [('name', profile['name']), ('contact', contact(profile))]
    if cv['summary']:
        blocks += [('heading', 'Professional Summary'), ('body', cv['summary'])]
    if cv['skills']:
        blocks.append(('heading', 'Skills'))
        blocks += [('body', group['category'] + ': ' + ', '.join(group['items'])) for group in cv['skills']]
    if cv['experience']:
        blocks.append(('heading', 'Work Experience'))
        for exp in cv['experience']:
            blocks += [('title', f"{exp['title']} | {exp['company']}"),
                       ('body', ' | '.join(filter(None, (exp['location'], date_range(exp['start'], exp['end'])))))]
            blocks += [('bullet', bullet) for bullet in exp['bullets']]
    if cv['projects']:
        blocks.append(('heading', 'Projects'))
        for project in cv['projects']:
            blocks.append(('title', project['name']))
            if project['tech']:
                blocks.append(('body', ', '.join(project['tech'])))
            if project.get('description'):
                blocks.append(('body', project['description']))
            blocks += [('bullet', bullet) for bullet in project['bullets']]
            if project['link']:
                blocks.append(('body', project['link']))
    if cv['education']:
        blocks.append(('heading', 'Education'))
        for edu in cv['education']:
            blocks += [('title', f"{edu['degree']} | {edu['institution']}"),
                       ('body', date_range(edu['start'], edu['end']))]
            if edu['details']:
                blocks.append(('body', edu['details']))
    if cv['certifications']:
        blocks.append(('heading', 'Certifications'))
        blocks += [('body', ' | '.join(filter(None, (c['name'], c['issuer'], c['year'])))) for c in cv['certifications']]
    return [(kind, text) for kind, text in blocks if text]


def cover_blocks(profile, letter, company, application_date=None):
    return [('name', profile['name']), ('contact', contact(profile)),
            ('body', application_date or date.today().strftime('%d %b %Y')), ('body', company),
            *[('letter', letter[k]) for k in ('greeting', 'opening', 'body_experience', 'body_skills', 'closing', 'sign_off')],
            ('body', profile['name'])]


def render_docx(blocks):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for kind, text in blocks:
        paragraph = doc.add_paragraph(style='List Bullet' if kind == 'bullet' else None)
        run = paragraph.add_run(text)
        if kind in ('name', 'heading', 'title'):
            run.bold = True
            paragraph.paragraph_format.keep_with_next = True
        if kind == 'name':
            run.font.size = Pt(17)
        elif kind == 'heading':
            run.font.size = Pt(12)
            paragraph.paragraph_format.space_before = Pt(8)
        elif kind == 'letter':
            paragraph.paragraph_format.space_after = Pt(10)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _pdf_fonts():
    # Linux substitutes Liberation Sans, a metrically compatible Arial font.
    roots = [Path('/usr/share/fonts/truetype/msttcorefonts'), Path('/usr/share/fonts/truetype/liberation2'),
             Path('/usr/share/fonts/truetype/liberation')]
    for root in roots:
        for regular, bold in [('Arial.ttf', 'Arial_Bold.ttf'), ('LiberationSans-Regular.ttf', 'LiberationSans-Bold.ttf')]:
            if (root / regular).exists() and (root / bold).exists():
                if 'CVArial' not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont('CVArial', str(root / regular)))
                    pdfmetrics.registerFont(TTFont('CVArialBold', str(root / bold)))
                return 'CVArial', 'CVArialBold'
    return 'Helvetica', 'Helvetica-Bold'


def render_pdf(blocks):
    regular, bold = _pdf_fonts()
    sizes = {'name': 17, 'heading': 12}
    styles = {kind: ParagraphStyle(kind, fontName=bold if kind in ('name', 'heading', 'title') else regular,
              fontSize=sizes.get(kind, 10.5), leading=sizes.get(kind, 10.5) * 1.15,
              spaceBefore=8 if kind == 'heading' else 0,
              spaceAfter=10 if kind == 'letter' else 4,
              leftIndent=12 if kind == 'bullet' else 0,
              bulletIndent=0, alignment=TA_LEFT, textColor=colors.black,
              keepWithNext=kind in ('name', 'heading', 'title'))
              for kind in ('name', 'contact', 'heading', 'title', 'body', 'bullet', 'letter')}
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=(612, 792), leftMargin=54, rightMargin=54,
                            topMargin=50.4, bottomMargin=50.4)
    flow = [Paragraph(escape(text).replace('\n', '<br/>'), styles[kind],
                      bulletText='•' if kind == 'bullet' else None) for kind, text in blocks]
    doc.build(flow)
    return stream.getvalue()


def page_count(pdf):
    return len(PdfReader(BytesIO(pdf)).pages)
