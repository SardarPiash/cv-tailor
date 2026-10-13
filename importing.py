from io import BytesIO
from docx import Document
from pypdf import PdfReader


def extract_text(content, filename):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError('Maximum upload size is 10 MB.')
    try:
        if filename.lower().endswith('.docx'):
            doc = Document(BytesIO(content))
            text = '\n'.join(p.text for p in doc.paragraphs)
            text += '\n' + '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        elif filename.lower().endswith('.pdf'):
            text = '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(content)).pages)
        else:
            raise ValueError('Upload a PDF or DOCX file.')
    except ValueError:
        raise
    except Exception:
        raise ValueError('Unable to read this file. Use an unencrypted PDF or valid DOCX.') from None
    if not text.strip():
        raise ValueError('No text found. Scanned PDFs need OCR first; upload a text PDF or DOCX.')
    if len(text) > 80000:
        raise ValueError('CV text is too long. Upload a shorter CV (under 80,000 characters).')
    return text
